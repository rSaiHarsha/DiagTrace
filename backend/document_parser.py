import io
import os
import json
from typing import Dict, Any, Optional

def parse_document_file(file_name: str, file_bytes: bytes, category: Optional[str] = None) -> str:
    """Extracts raw text content, layout, tables, and SysML diagrams from uploaded document files (.txt, .pdf, .docx, .csv, .json, .log, .md, .png, .jpg, .jpeg, .webp)."""
    ext = os.path.splitext(file_name)[1].lower()
    
    if ext in ['.png', '.jpg', '.jpeg', '.webp', '.bmp', '.tiff']:
        return extract_image_text(file_bytes, file_name, category)
    elif ext in ['.pdf']:
        return extract_pdf_text(file_bytes, category=category)
    elif ext in ['.docx']:
        return extract_docx_text(file_bytes)
    elif ext in ['.json']:
        return extract_json_text(file_bytes)
    else:
        # Default text-based formats (.txt, .md, .log, .csv)
        return extract_plain_text(file_bytes)

def extract_plain_text(file_bytes: bytes) -> str:
    """Decodes plain text files with encoding fallbacks."""
    for encoding in ['utf-8', 'utf-8-sig', 'latin-1', 'cp1252']:
        try:
            text = file_bytes.decode(encoding)
            return text.strip()
        except UnicodeDecodeError:
            continue
    return file_bytes.decode('utf-8', errors='ignore').strip()

import base64

def extract_image_text(file_bytes: bytes, file_name: str, category: Optional[str] = None) -> str:
    """Extracts text and converts ECU Architecture diagrams to formal SysML (v2/PlantUML) code using NVIDIA Llama 3.2 Vision."""
    b64_img = base64.b64encode(file_bytes).decode("utf-8")
    try:
        from backend.nvidia_client import query_nvidia_vision_sysml, query_nvidia_vision_ocr
        if category and "architecture" in category.lower():
            print(f"[SYSML_CONVERTER] Processing ECU Architecture diagram image '{file_name}' with SysML vision converter...")
            return query_nvidia_vision_sysml(b64_img)
        else:
            return query_nvidia_vision_ocr(b64_img)
    except Exception as e:
        print(f"Error parsing image '{file_name}' with Vision OCR: {e}")
        return f"[Image Document: {file_name}]\n(Image content could not be converted via Vision API: {e})"

def extract_pdf_text(file_bytes: bytes, category: Optional[str] = None) -> str:
    """Extracts layout-structured text, headings, Markdown tables, and NVIDIA Vision OCR/SysML diagrams from PDF files using PyMuPDF."""
    try:
        import pymupdf
        doc = pymupdf.open(stream=file_bytes, filetype="pdf")
        pages_output = []
        
        for i, page in enumerate(doc):
            page_text_blocks = []
            
            # 1. Extract structural text blocks (preserves headings, paragraphs, layout order)
            try:
                blocks = page.get_text("blocks")
                # blocks format: (x0, y0, x1, y1, text, block_no, block_type)
                for b in blocks:
                    block_txt = b[4].strip() if len(b) > 4 and isinstance(b[4], str) else ""
                    if block_txt:
                        page_text_blocks.append(block_txt)
            except Exception:
                txt = page.get_text("text")
                if txt and txt.strip():
                    page_text_blocks.append(txt.strip())

            # 2. Extract tabular data using PyMuPDF table finder if available
            try:
                tabs = page.find_tables()
                if tabs and tabs.tables:
                    for tab in tabs.tables:
                        df_tab = tab.extract()
                        if df_tab:
                            md_rows = []
                            for r_idx, row in enumerate(df_tab):
                                cleaned_row = [str(cell).strip() if cell else "" for cell in row]
                                md_rows.append("| " + " | ".join(cleaned_row) + " |")
                                if r_idx == 0:
                                    md_rows.append("| " + " | ".join(["---"] * len(cleaned_row)) + " |")
                            if md_rows:
                                page_text_blocks.append("\n".join(md_rows))
            except Exception as tab_err:
                print(f"PyMuPDF Table Extraction notice for page {i+1}: {tab_err}")

            combined_page_text = "\n\n".join(page_text_blocks).strip()

            # 3. Check if page is scanned/image-heavy or if Architecture category is selected
            is_architecture_category = bool(category and "architecture" in category.lower())
            if len(combined_page_text) < 40 or is_architecture_category:
                try:
                    from backend.nvidia_client import query_nvidia_vision_ocr, query_nvidia_vision_sysml
                    pix = page.get_pixmap(dpi=150)
                    img_bytes = pix.tobytes("png")
                    b64_img = base64.b64encode(img_bytes).decode("utf-8")
                    
                    if is_architecture_category:
                        print(f"[SYSML_CONVERTER] Page {i+1} ECU Architecture Diagram detected. Triggering NVIDIA Llama 3.2 Vision SysML Converter...")
                        vision_text = query_nvidia_vision_sysml(b64_img)
                        if vision_text:
                            combined_page_text = f"{combined_page_text}\n\n[NVIDIA Vision SysML Diagram Extraction]\n{vision_text}".strip()
                    else:
                        print(f"[VISION_OCR] Page {i+1} appears image-heavy/scanned. Triggering NVIDIA Llama 3.2 Vision OCR...")
                        ocr_text = query_nvidia_vision_ocr(b64_img)
                        if ocr_text:
                            combined_page_text = f"[NVIDIA Vision OCR Extracted Content]\n{ocr_text}"
                except Exception as ocr_err:
                    print(f"NVIDIA Vision OCR fallback notice for page {i+1}: {ocr_err}")

            if combined_page_text:
                pages_output.append(f"--- Page {i+1} ---\n{combined_page_text}")
                
        doc.close()
        if pages_output:
            return "\n\n".join(pages_output)
    except Exception as e:
        print(f"Error parsing PDF with PyMuPDF: {e}")

    # Legacy pypdf fallback if PyMuPDF fails
    try:
        from pypdf import PdfReader
        pdf_file = io.BytesIO(file_bytes)
        reader = PdfReader(pdf_file)
        pages_text = []
        for i, page in enumerate(reader.pages):
            txt = page.extract_text()
            if txt and txt.strip():
                pages_text.append(f"--- Page {i+1} ---\n{txt.strip()}")
        if pages_text:
            return "\n\n".join(pages_text)
    except Exception:
        pass
        
    return extract_plain_text(file_bytes)

def extract_docx_text(file_bytes: bytes) -> str:
    """Extracts paragraphs and tables from Word (.docx) files using python-docx."""
    try:
        import docx
        docx_file = io.BytesIO(file_bytes)
        doc = docx.Document(docx_file)
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
        
        # Extract tables if present
        for table in doc.tables:
            for row in table.rows:
                row_txt = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                if row_txt:
                    paragraphs.append(row_txt)
                    
        if paragraphs:
            return "\n\n".join(paragraphs)
    except Exception as e:
        print(f"Error parsing DOCX with python-docx: {e}")
        
    return extract_plain_text(file_bytes)

def extract_json_text(file_bytes: bytes) -> str:
    """Formats JSON files into clean readable markdown/text representation."""
    try:
        text = extract_plain_text(file_bytes)
        data = json.loads(text)
        return json.dumps(data, indent=2)
    except Exception:
        return extract_plain_text(file_bytes)
